#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""从 QQ 收到的文档里抽取可读文本：pdf / xlsx / xlsm / pptx。

设计口径（2026-10-28 主人拍板）：
- 文本优先：PDF 先取文本层（PyMuPDF 对中文 CID 字体比 pdftotext 好得多），
  xlsx 用 openpyxl 读单元格，pptx 直接读 slide XML；
- 视觉兜底：PDF 某页几乎没字（扫描件/整页图片）才渲染成 PNG，交给桥的识图模型逐字转录；
- 不把整份文档无脑转图片：表格数字看图会看错，且每页一次视觉调用又慢又贵。

输出 UTF-8 JSON 到 stdout（Node 侧 execFile 读）。任何异常都返回 {ok:false,error}，不抛栈。
"""
import sys
import os
import io
import json
import base64
import html
import re
import zipfile
import argparse

# Windows 控制台默认 GBK：必须显式 UTF-8，否则 Node 收到乱码/直接抛 UnicodeEncodeError
try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

MIN_PAGE_CHARS = 30          # 少于此字数视为"这页没文本层"，需要视觉
MAX_PDF_PAGES = 40           # 最多读多少页
MAX_RENDER_IMAGES = 8        # 最多给识图模型多少张页图
RENDER_DPI = 150


def _ok(kind, pages, meta=None, note=None):
    return {"ok": True, "kind": kind, "pages": pages, "meta": meta or {}, "note": note}


def pdf_extract(path, max_chars):
    import pymupdf
    doc = pymupdf.open(path)
    pages = []
    total = 0
    images = 0
    for i, page in enumerate(doc):
        if i >= MAX_PDF_PAGES:
            break
        text = (page.get_text("text") or "").strip()
        chars = len(text)
        item = {"index": i + 1, "text": text, "chars": chars, "needsVision": False}
        if chars < MIN_PAGE_CHARS and images < MAX_RENDER_IMAGES:
            try:
                pix = page.get_pixmap(dpi=RENDER_DPI, colorspace=pymupdf.csRGB)
                item["image"] = base64.b64encode(pix.tobytes("png")).decode("ascii")
                item["needsVision"] = True
                images += 1
            except Exception as exc:  # 渲染失败不影响文字页
                item["note"] = f"页面渲染失败：{exc}"
        total += chars
        pages.append(item)
        if total >= max_chars:
            break
    note = None
    if doc.page_count > len(pages):
        note = f"只读了前 {len(pages)} 页（共 {doc.page_count} 页）"
    return _ok("pdf", pages, {"pageCount": doc.page_count, "rendered": images}, note)


def xlsx_extract(path, max_chars):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    pages = []
    total = 0
    for ws in wb.worksheets[:12]:
        max_row = min(int(ws.max_row or 0), 400)
        max_col = min(int(ws.max_column or 0), 40)
        lines = [f"# 工作表：{ws.title}（{max_row} 行 × {max_col} 列）"]
        used = 0
        for r, row in enumerate(
            ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col, values_only=True), 1
        ):
            cells = []
            for c in row:
                if c is None:
                    cells.append("")
                else:
                    s = str(c).replace("\r", " ").replace("\n", " ").strip()
                    cells.append(s[:200])
            if any(cells):
                line = f"{r}\t" + "\t".join(cells).rstrip()
                lines.append(line)
                used += len(line)
            if used >= max_chars:
                lines.append(f"…（本表超过 {max_chars} 字符，已截断）")
                break
        text = "\n".join(lines)
        pages.append({"index": len(pages) + 1, "name": ws.title, "text": text, "chars": len(text), "needsVision": False})
        total += len(text)
        if total >= max_chars:
            break
    note = "只读单元格与缓存值；图表/图片/批注不可读，公式显示计算结果"
    if len(wb.sheetnames) > len(pages):
        note += f"；只读了前 {len(pages)} 个工作表（共 {len(wb.sheetnames)} 个）"
    return _ok("xlsx", pages, {"sheets": len(wb.sheetnames)}, note)


def _text_from_xml(xml):
    texts = re.findall(r"<a:t>(.*?)</a:t>", xml, re.S)
    return "\n".join(t.strip() for t in (html.unescape(x) for x in texts) if t.strip())


def pptx_extract(path, max_chars):
    z = zipfile.ZipFile(path)
    slide_names = [n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)]
    slide_names.sort(key=lambda n: int(re.search(r"slide(\d+)\.xml$", n).group(1)))
    pages = []
    total = 0
    for i, name in enumerate(slide_names[:30], 1):
        xml = z.read(name).decode("utf-8", "replace")
        text = _text_from_xml(xml)
        # 演讲者备注也是内容的一部分
        note_name = f"ppt/notesSlides/notesSlide{i}.xml"
        if note_name in z.namelist():
            notes = _text_from_xml(z.read(note_name).decode("utf-8", "replace"))
            if notes:
                text += "\n[备注]\n" + notes
        text = text[:max_chars]
        pages.append({"index": i, "name": f"第{i}张幻灯片", "text": text, "chars": len(text), "needsVision": False})
        total += len(text)
        if total >= max_chars:
            break
    note = "读的是文字层（标题/正文/表格/备注）；图表、SmartArt 图形与版式不可读"
    if len(slide_names) > len(pages):
        note += f"；只读了前 {len(pages)} 张（共 {len(slide_names)} 张）"
    return _ok("pptx", pages, {"slides": len(slide_names)}, note)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--path", required=True)
    ap.add_argument("--max-chars", type=int, default=30000)
    args = ap.parse_args()
    path = args.path
    ext = os.path.splitext(path)[1].lower()
    max_chars = max(2000, min(60000, args.max_chars))
    try:
        if ext == ".pdf":
            out = pdf_extract(path, max_chars)
        elif ext in (".xlsx", ".xlsm"):
            out = xlsx_extract(path, max_chars)
        elif ext in (".pptx", ".pptm"):
            out = pptx_extract(path, max_chars)
        else:
            out = {"ok": False, "error": f"不支持的文档类型：{ext}"}
    except Exception as exc:
        out = {"ok": False, "error": f"{type(exc).__name__}: {exc}"}
    json.dump(out, sys.stdout, ensure_ascii=False)
    sys.stdout.write("\n")


if __name__ == "__main__":
    main()
